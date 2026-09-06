"""Export tests — is every answer reachable offline?  (tests/test_export.py)

export.py freezes engine output into site/data/*.json for an app that has
no Python. These check completeness (a file per entity), parity (the JSON
says what the engine says) and hygiene (valid JSON, no stale files).
"""

from __future__ import annotations

import json
import sqlite3

import pytest

from src.engine import Engine
from src.export import Exporter, write_audit_excel

DB = "build/wine.db"


@pytest.fixture(scope="module")
def engine() -> Engine:
    return Engine(DB)


@pytest.fixture
def exported(engine, tmp_path):
    """A full export into a temp dir — never touches the real site/."""
    Exporter(engine, tmp_path).run()
    return tmp_path


@pytest.fixture(scope="module")
def con():
    c = sqlite3.connect(f"file:{DB}?mode=ro", uri=True)
    c.row_factory = sqlite3.Row
    return c


def test_writes_a_file_per_entity(exported, con):
    """Every wine, tag and region must be reachable offline."""
    for table, folder in (
        ("wines", "wines"),
        ("food_tags", "pairings"),
        ("regions", "regions"),
    ):
        for row in con.execute(f"SELECT id FROM {table}"):
            assert (
                exported / folder / f"{row['id']}.json"
            ).exists(), f"missing {folder}/{row['id']}.json"


def test_top_level_files_exist(exported):
    for name in ("index.json", "search-docs.json", "flavours.json"):
        assert (exported / name).exists()


def test_everything_is_valid_json(exported):
    for path in exported.rglob("*.json"):
        json.loads(path.read_text(encoding="utf-8"))


def test_pairings_match_the_engine(exported, engine):
    """The site and the CLI must never disagree (README §5.3 parity)."""
    for path in (exported / "pairings").glob("*.json"):
        payload = json.loads(path.read_text(encoding="utf-8"))
        live = engine.pair_food([payload["tag"]], limit=50)
        assert [w.wine_id for w in live] == [w["wine_id"] for w in payload["wines"]]


def test_pairings_carry_whys_and_avoid_lists(exported):
    """app.js needs both to do the client-side merge."""
    for path in (exported / "pairings").glob("*.json"):
        payload = json.loads(path.read_text(encoding="utf-8"))
        assert "avoid" in payload
        for wine in payload["wines"]:
            assert wine["why"] and all(w.strip() for w in wine["why"])


def test_search_docs_cover_every_kind(exported, con):
    docs = json.loads((exported / "search-docs.json").read_text(encoding="utf-8"))
    kinds = {d["kind"] for d in docs}
    expected = {"wine", "food", "region"}
    present = {
        k
        for k in expected
        if con.execute(
            f"SELECT COUNT(*) FROM "
            f"{'wines' if k == 'wine' else k + 's' if k == 'region' else 'food_tags'}"
        ).fetchone()[0]
    }
    assert present <= kinds


def test_search_docs_are_normalized(exported):
    """Index-time and query-time rules must match, or accented names miss."""
    from src.engine import normalise

    docs = json.loads((exported / "search-docs.json").read_text(encoding="utf-8"))
    for doc in docs:
        for field, value in doc.items():
            if field in ("id", "kind", "ref", "title"):
                continue
            assert value == normalise(value), f"{doc['id']}.{field} not normalized"


def test_index_counts_match_the_database(exported, con):
    index = json.loads((exported / "index.json").read_text(encoding="utf-8"))
    for key, table in (
        ("wines", "wines"),
        ("food_tags", "food_tags"),
        ("regions", "regions"),
    ):
        assert (
            index["counts"][key]
            == con.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        )


def test_stale_files_are_cleared(engine, tmp_path):
    """A deleted wine must not linger on the phone."""
    stale = tmp_path / "wines" / "chateau-deleted.json"
    stale.parent.mkdir(parents=True)
    stale.write_text("{}")
    Exporter(engine, tmp_path).run()
    assert not stale.exists()


def test_audit_xlsx_has_a_row_per_wine(engine, con, tmp_path):
    from openpyxl import load_workbook

    path = write_audit_excel(engine, tmp_path / "wine_audit.xlsx")
    ws = load_workbook(path).active
    count = con.execute("SELECT COUNT(*) FROM wines").fetchone()[0]
    assert ws.max_row == count + 1  # +1 for the header
    assert ws["A1"].value == "id"


def test_audit_columns_are_in_schema_order(engine, tmp_path):
    """Regression: a set-valued SCALE_DIMS shuffled these between builds."""
    from openpyxl import load_workbook
    from src.schema import SCALE_DIMS

    ws = load_workbook(write_audit_excel(engine, tmp_path / "a.xlsx")).active
    headers = [c.value for c in ws[1]]
    assert headers[4 : 4 + len(SCALE_DIMS)] == list(SCALE_DIMS)
